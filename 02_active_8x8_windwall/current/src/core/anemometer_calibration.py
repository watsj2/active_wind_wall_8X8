"""
Anemometer/tach calibration helpers for uniform-flow surveys.

The GUI collects raw per-motor telemetry samples. This module turns those
samples into motor, row, column, heatmap, and compensation tables, then writes
them to an Excel workbook.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

import numpy as np

from config import (
    GRID_COLS,
    GRID_ROWS,
    NUM_MOTORS,
    PWM_MAX,
    PWM_MIN,
    TACH_PULSES_PER_REV,
    WALL_MOTOR_GRID,
)


MOTOR_GRID_POSITION: dict[int, tuple[int, int]] = {
    motor_id: (row, col)
    for row, row_motors in enumerate(WALL_MOTOR_GRID)
    for col, motor_id in enumerate(row_motors)
}


@dataclass(frozen=True)
class AnemometerSample:
    """One raw telemetry snapshot for all motors."""

    timestamp: str
    elapsed_s: float
    tach_hz: np.ndarray


@dataclass(frozen=True)
class AnemometerCalibrationResult:
    """Computed calibration products for export and later compensation."""

    commanded_pwm_us: float
    target_hz: float
    target_rpm: float
    sample_count: int
    motor_summary: list[dict[str, Any]]
    row_summary: list[dict[str, Any]]
    column_summary: list[dict[str, Any]]
    mean_hz_by_motor: np.ndarray
    suggested_pwm_by_motor: np.ndarray
    offset_us_by_motor: np.ndarray
    mean_hz_grid: np.ndarray
    suggested_pwm_grid: np.ndarray
    offset_us_grid: np.ndarray


@dataclass(frozen=True)
class TestoMeasurement:
    """One Testo 440 dP CSV summarized at one tunnel grid point."""

    source_path: str
    row: int
    col: int
    distance_m: float
    start_time: datetime | None
    end_time: datetime | None
    duration_s: float
    sample_count: int
    mean_ms: float
    median_ms: float
    std_ms: float
    min_ms: float
    max_ms: float
    pressure_hpa_mean: float


@dataclass(frozen=True)
class TestoCorrelationResult:
    """Merged Testo point data and GUI log data on the tunnel measurement grid."""

    distance_m: float
    log_path: str | None
    testo_total_duration_s: float
    testo_mean_duration_s: float
    gui_matched_duration_s: float
    gui_mean_pwm_us: float
    gui_min_pwm_us: float
    gui_max_pwm_us: float
    gui_log_sample_count: int
    matched_log_count: int
    measurements: list[TestoMeasurement]
    testo_grid_ms: np.ndarray
    expanded_speed_grid_ms: np.ndarray
    gui_pwm_grid: np.ndarray
    gui_tach_hz_grid: np.ndarray
    gui_rpm_grid: np.ndarray
    comparison_rows: list[dict[str, Any]]
    correlations: dict[str, float]


def _finite(values: np.ndarray) -> np.ndarray:
    return values[np.isfinite(values)]


def _stat(values: np.ndarray, reducer, default: float = float("nan")) -> float:
    finite = _finite(values)
    if finite.size == 0:
        return default
    return float(reducer(finite))


def _motor_values_to_grid(values: np.ndarray) -> np.ndarray:
    grid = np.full((GRID_ROWS, GRID_COLS), np.nan, dtype=np.float64)
    for motor_id, (row, col) in MOTOR_GRID_POSITION.items():
        grid[row, col] = values[motor_id]
    return grid


def _measurement_block_indices(row: int, col: int, rows: int = 4, cols: int = 4) -> tuple[np.ndarray, np.ndarray]:
    """Return motor-grid row/col indices covered by one Testo measurement point."""
    row_blocks = np.array_split(np.arange(GRID_ROWS), rows)
    col_blocks = np.array_split(np.arange(GRID_COLS), cols)
    return row_blocks[row], col_blocks[col]


def _block_average_grid(values_by_motor: np.ndarray, rows: int = 4, cols: int = 4) -> np.ndarray:
    """Average 8x8 motor values into the 4x4 Testo measurement grid."""
    motor_grid = _motor_values_to_grid(np.asarray(values_by_motor, dtype=np.float64))
    block_grid = np.full((rows, cols), np.nan, dtype=np.float64)
    for row in range(rows):
        motor_rows, _ = _measurement_block_indices(row, 0, rows, cols)
        for col in range(cols):
            _, motor_cols = _measurement_block_indices(0, col, rows, cols)
            block = motor_grid[np.ix_(motor_rows, motor_cols)]
            finite = _finite(block)
            if finite.size:
                block_grid[row, col] = float(np.mean(finite))
    return block_grid


def _expand_measurement_grid_to_motor_grid(measurement_grid: np.ndarray) -> np.ndarray:
    """Expand a 4x4 Testo grid to the 8x8 motor grid by filling each 2x2 block."""
    measurement_grid = np.asarray(measurement_grid, dtype=np.float64)
    expanded = np.full((GRID_ROWS, GRID_COLS), np.nan, dtype=np.float64)
    rows, cols = measurement_grid.shape
    for row in range(rows):
        motor_rows, _ = _measurement_block_indices(row, 0, rows, cols)
        for col in range(cols):
            _, motor_cols = _measurement_block_indices(0, col, rows, cols)
            expanded[np.ix_(motor_rows, motor_cols)] = measurement_grid[row, col]
    return expanded


def _parse_float(value: Any) -> float:
    text = str(value).strip()
    if not text:
        return float("nan")
    try:
        return float(text)
    except ValueError:
        return float(text.replace(",", "."))


def _parse_testo_datetime(date_value: Any, time_value: Any) -> datetime | None:
    text = f"{str(date_value).strip()} {str(time_value).strip()}"
    if not text.strip():
        return None
    for fmt in ("%m/%d/%Y %H:%M:%S", "%m/%d/%y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _format_datetime(value: datetime | None) -> str | None:
    return value.isoformat(sep=" ") if value is not None else None


def _parse_testo_position(path: str | Path) -> tuple[int, int]:
    match = re.search(r"R(\d+)C(\d+)", Path(path).stem.upper())
    if not match:
        raise ValueError(f"Filename must contain R#C#: {path}")
    row = int(match.group(1)) - 1
    col = int(match.group(2)) - 1
    if row < 0 or col < 0 or row >= 4 or col >= 4:
        raise ValueError(f"Testo point must be within R1C1..R4C4: {path}")
    return row, col


def parse_testo_csv(path: str | Path, distance_m: float) -> TestoMeasurement:
    """Parse one Testo 440 dP logger CSV and summarize the velocity column."""
    csv_path = Path(path)
    row_index, col_index = _parse_testo_position(csv_path)
    with csv_path.open("r", encoding="utf-8-sig", newline=None) as fh:
        rows = list(csv.reader(fh))

    header_index = None
    speed_col = None
    pressure_col = None
    date_col = None
    time_col = None
    for idx, row in enumerate(rows):
        labels = [cell.strip() for cell in row]
        if labels and labels[0] == "Measured Points":
            header_index = idx
            for col, label in enumerate(labels):
                label_lower = label.lower()
                if label_lower == "date":
                    date_col = col
                if label_lower == "time":
                    time_col = col
                if "m/s" in label_lower and speed_col is None:
                    speed_col = col
                if "hpa" in label_lower and pressure_col is None:
                    pressure_col = col
            break

    if header_index is None or speed_col is None:
        raise ValueError(f"Could not find Testo velocity data in {csv_path}")

    speed_values: list[float] = []
    pressure_values: list[float] = []
    timestamps: list[datetime] = []
    for row in rows[header_index + 1:]:
        if not row or not row[0].strip().isdigit():
            continue
        if date_col is not None and time_col is not None and len(row) > max(date_col, time_col):
            timestamp = _parse_testo_datetime(row[date_col], row[time_col])
            if timestamp is not None:
                timestamps.append(timestamp)
        if len(row) > speed_col:
            value = _parse_float(row[speed_col])
            if np.isfinite(value):
                speed_values.append(value)
        if pressure_col is not None and len(row) > pressure_col:
            value = _parse_float(row[pressure_col])
            if np.isfinite(value):
                pressure_values.append(value)

    if not speed_values:
        raise ValueError(f"No finite Testo velocity samples found in {csv_path}")

    speeds = np.asarray(speed_values, dtype=np.float64)
    pressures = np.asarray(pressure_values, dtype=np.float64)
    start_time = timestamps[0] if timestamps else None
    end_time = timestamps[-1] if timestamps else None
    duration_s = (
        float((end_time - start_time).total_seconds())
        if start_time is not None and end_time is not None
        else float("nan")
    )
    return TestoMeasurement(
        source_path=str(csv_path),
        row=row_index,
        col=col_index,
        distance_m=float(distance_m),
        start_time=start_time,
        end_time=end_time,
        duration_s=duration_s,
        sample_count=int(speeds.size),
        mean_ms=float(np.mean(speeds)),
        median_ms=float(np.median(speeds)),
        std_ms=float(np.std(speeds)),
        min_ms=float(np.min(speeds)),
        max_ms=float(np.max(speeds)),
        pressure_hpa_mean=float(np.mean(pressures)) if pressures.size else float("nan"),
    )


def _empty_gui_log_data() -> tuple[dict[str, np.ndarray], dict[str, Any]]:
    nan_values = np.full(NUM_MOTORS, np.nan, dtype=np.float64)
    return (
        {"pwm": nan_values.copy(), "tach_hz": nan_values.copy(), "rpm": nan_values.copy()},
        {
            "start_time": None,
            "end_time": None,
            "duration_s": float("nan"),
            "gui_mean_pwm_us": float("nan"),
            "gui_min_pwm_us": float("nan"),
            "gui_max_pwm_us": float("nan"),
            "gui_pwm_value_count": 0,
            "gui_log_sample_count": 0,
        },
    )


def _parse_iso_timestamp(value: Any) -> datetime | None:
    try:
        return datetime.fromisoformat(str(value).strip())
    except ValueError:
        return None


def _read_gui_log_data(
    log_path: str | Path | None,
    start_time: datetime | None = None,
    end_time: datetime | None = None,
) -> tuple[dict[str, np.ndarray], dict[str, Any]]:
    """Return mean PWM/tach/RPM arrays and metadata from a GUI log time window."""
    if not log_path:
        return _empty_gui_log_data()

    path = Path(log_path)
    if not path.exists():
        return _empty_gui_log_data()

    sums = {
        "pwm": np.zeros(NUM_MOTORS, dtype=np.float64),
        "tach_hz": np.zeros(NUM_MOTORS, dtype=np.float64),
        "rpm": np.zeros(NUM_MOTORS, dtype=np.float64),
    }
    counts = {key: np.zeros(NUM_MOTORS, dtype=np.int64) for key in sums}
    first_timestamp: datetime | None = None
    last_timestamp: datetime | None = None
    log_sample_count = 0
    pwm_total = 0.0
    pwm_count = 0
    pwm_min = float("nan")
    pwm_max = float("nan")

    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        try:
            header = next(reader)
        except StopIteration:
            return _empty_gui_log_data()

        index_by_name = {name: idx for idx, name in enumerate(header)}
        timestamp_index = index_by_name.get("timestamp", 0)
        column_indexes = {
            key: [index_by_name.get(f"{key}_{motor_id}") for motor_id in range(NUM_MOTORS)]
            for key in sums
        }

        for line in reader:
            timestamp = None
            if timestamp_index < len(line):
                timestamp = _parse_iso_timestamp(line[timestamp_index])
            if start_time is not None and end_time is not None:
                if timestamp is None or timestamp < start_time or timestamp > end_time:
                    continue

            log_sample_count += 1
            if timestamp is not None:
                if first_timestamp is None:
                    first_timestamp = timestamp
                last_timestamp = timestamp
            for key, indexes in column_indexes.items():
                for motor_id, column_index in enumerate(indexes):
                    if column_index is None or column_index >= len(line):
                        continue
                    try:
                        value = float(line[column_index])
                    except ValueError:
                        continue
                    if np.isfinite(value):
                        sums[key][motor_id] += value
                        counts[key][motor_id] += 1
                        if key == "pwm":
                            pwm_total += value
                            pwm_count += 1
                            pwm_min = value if not np.isfinite(pwm_min) else min(pwm_min, value)
                            pwm_max = value if not np.isfinite(pwm_max) else max(pwm_max, value)

    means: dict[str, np.ndarray] = {}
    for key in sums:
        means[key] = np.full(NUM_MOTORS, np.nan, dtype=np.float64)
        valid = counts[key] > 0
        means[key][valid] = sums[key][valid] / counts[key][valid]

    metadata = {
        "start_time": first_timestamp,
        "end_time": last_timestamp,
        "duration_s": (
            float((last_timestamp - first_timestamp).total_seconds())
            if first_timestamp is not None and last_timestamp is not None
            else float("nan")
        ),
        "gui_mean_pwm_us": float(pwm_total / pwm_count) if pwm_count else float("nan"),
        "gui_min_pwm_us": float(pwm_min),
        "gui_max_pwm_us": float(pwm_max),
        "gui_pwm_value_count": int(pwm_count),
        "gui_log_sample_count": int(log_sample_count),
    }
    return means, metadata


def _candidate_log_paths(log_path: str | Path | None) -> list[Path]:
    if log_path:
        return [Path(log_path)]
    return sorted(Path("logs").glob("flight_log_*.csv"), key=lambda path: path.stat().st_mtime, reverse=True)


def _read_gui_log_summary(log_path: str | Path) -> dict[str, Any]:
    """Read only timestamp range metadata from a GUI flight log."""
    first_timestamp: datetime | None = None
    last_timestamp: datetime | None = None
    row_count = 0
    path = Path(log_path)

    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        try:
            header = next(reader)
        except StopIteration:
            return {
                "start_time": None,
                "end_time": None,
                "duration_s": float("nan"),
                "gui_log_sample_count": 0,
            }
        timestamp_index = {name: idx for idx, name in enumerate(header)}.get("timestamp", 0)
        for row in reader:
            row_count += 1
            if timestamp_index >= len(row):
                continue
            timestamp = _parse_iso_timestamp(row[timestamp_index])
            if timestamp is None:
                continue
            if first_timestamp is None:
                first_timestamp = timestamp
            last_timestamp = timestamp

    return {
        "start_time": first_timestamp,
        "end_time": last_timestamp,
        "duration_s": (
            float((last_timestamp - first_timestamp).total_seconds())
            if first_timestamp is not None and last_timestamp is not None
            else float("nan")
        ),
        "gui_log_sample_count": int(row_count),
    }


def _overlap_seconds(start_a: datetime, end_a: datetime, start_b: datetime, end_b: datetime) -> float:
    start = max(start_a, start_b)
    end = min(end_a, end_b)
    return max(0.0, float((end - start).total_seconds()))


def _aligned_testo_window_for_log(
    measurement: TestoMeasurement,
    log_start: datetime | None,
    log_end: datetime | None,
) -> tuple[datetime | None, datetime | None, float]:
    """Align a Testo clock window to a GUI log date and return overlap seconds.

    Testo meters often have the wrong calendar date even when the time-of-day is
    set correctly, so matching is based on the Testo time-of-day and the GUI log
    date. This still respects midnight rollover and the Testo sample duration.
    """
    if measurement.start_time is None or measurement.end_time is None or log_start is None or log_end is None:
        return None, None, 0.0

    duration_s = (measurement.end_time - measurement.start_time).total_seconds()
    if duration_s < 0:
        duration_s += 24 * 60 * 60
    duration = timedelta(seconds=max(0.0, duration_s))

    base_dates = {
        log_start.date() + timedelta(days=offset)
        for offset in (-1, 0, 1)
    } | {
        log_end.date() + timedelta(days=offset)
        for offset in (-1, 0, 1)
    }

    best_start = None
    best_end = None
    best_overlap = 0.0
    for base_date in base_dates:
        aligned_start = datetime.combine(base_date, measurement.start_time.time())
        aligned_end = aligned_start + duration
        overlap = _overlap_seconds(aligned_start, aligned_end, log_start, log_end)
        if overlap > best_overlap:
            best_start = aligned_start
            best_end = aligned_end
            best_overlap = overlap

    return best_start, best_end, best_overlap


def _load_log_summaries(log_paths: list[Path]) -> list[dict[str, Any]]:
    summaries: list[dict[str, Any]] = []
    for path in log_paths:
        metadata = _read_gui_log_summary(path)
        if metadata["gui_log_sample_count"]:
            summaries.append({"path": path, **metadata})
    return summaries


def _best_log_match(
    measurement: TestoMeasurement,
    log_summaries: list[dict[str, Any]],
) -> tuple[Path | None, datetime | None, datetime | None, float]:
    best_path = None
    best_start = None
    best_end = None
    best_overlap = 0.0
    for summary in log_summaries:
        aligned_start, aligned_end, overlap = _aligned_testo_window_for_log(
            measurement,
            summary["start_time"],
            summary["end_time"],
        )
        if overlap > best_overlap:
            best_path = summary["path"]
            best_start = aligned_start
            best_end = aligned_end
            best_overlap = overlap

    if best_path is None and len(log_summaries) == 1 and measurement.start_time is None:
        summary = log_summaries[0]
        return summary["path"], summary["start_time"], summary["end_time"], summary["duration_s"]
    return best_path, best_start, best_end, best_overlap


def _pearson_from_rows(rows: list[dict[str, Any]], field: str) -> float:
    speed = np.asarray([row["testo_mean_ms"] for row in rows], dtype=np.float64)
    other = np.asarray([row[field] for row in rows], dtype=np.float64)
    valid = np.isfinite(speed) & np.isfinite(other)
    if np.count_nonzero(valid) < 2:
        return float("nan")
    if np.std(speed[valid]) == 0 or np.std(other[valid]) == 0:
        return float("nan")
    return float(np.corrcoef(speed[valid], other[valid])[0, 1])


def create_testo_correlation_result(
    csv_paths: list[str | Path],
    distance_m: float,
    log_path: str | Path | None = None,
) -> TestoCorrelationResult:
    """Merge Testo point measurements with block-averaged GUI flight-log data."""
    measurements = [parse_testo_csv(path, distance_m) for path in csv_paths]
    if not measurements:
        raise ValueError("No Testo CSV files were selected")

    testo_grid = np.full((4, 4), np.nan, dtype=np.float64)
    for measurement in measurements:
        testo_grid[measurement.row, measurement.col] = measurement.mean_ms

    log_summaries = _load_log_summaries(_candidate_log_paths(log_path))
    gui_pwm_grid = np.full((4, 4), np.nan, dtype=np.float64)
    gui_tach_grid = np.full((4, 4), np.nan, dtype=np.float64)
    gui_rpm_grid = np.full((4, 4), np.nan, dtype=np.float64)

    comparison_rows: list[dict[str, Any]] = []
    matched_logs: set[str] = set()
    gui_matched_duration_s = 0.0
    gui_log_sample_count = 0
    gui_pwm_total = 0.0
    gui_pwm_count = 0
    gui_min_pwm = float("nan")
    gui_max_pwm = float("nan")
    for measurement in sorted(measurements, key=lambda item: (item.row, item.col)):
        row = measurement.row
        col = measurement.col
        match_path, match_start, match_end, match_overlap_s = _best_log_match(measurement, log_summaries)
        gui_metadata = _empty_gui_log_data()[1]

        if match_path is not None and match_start is not None and match_end is not None:
            gui_means, gui_metadata = _read_gui_log_data(match_path, match_start, match_end)
            point_pwm_grid = _block_average_grid(gui_means["pwm"])
            point_tach_grid = _block_average_grid(gui_means["tach_hz"])
            point_rpm_grid = _block_average_grid(gui_means["rpm"])
            gui_pwm_grid[row, col] = point_pwm_grid[row, col]
            gui_tach_grid[row, col] = point_tach_grid[row, col]
            gui_rpm_grid[row, col] = point_rpm_grid[row, col]
            matched_logs.add(match_path.name)

            if np.isfinite(gui_metadata["duration_s"]):
                gui_matched_duration_s += float(gui_metadata["duration_s"])
            gui_log_sample_count += int(gui_metadata["gui_log_sample_count"])
            if gui_metadata["gui_pwm_value_count"] and np.isfinite(gui_metadata["gui_mean_pwm_us"]):
                gui_pwm_total += float(gui_metadata["gui_mean_pwm_us"]) * int(gui_metadata["gui_pwm_value_count"])
                gui_pwm_count += int(gui_metadata["gui_pwm_value_count"])
            if np.isfinite(gui_metadata["gui_min_pwm_us"]):
                gui_min_pwm = (
                    float(gui_metadata["gui_min_pwm_us"])
                    if not np.isfinite(gui_min_pwm)
                    else min(gui_min_pwm, float(gui_metadata["gui_min_pwm_us"]))
                )
            if np.isfinite(gui_metadata["gui_max_pwm_us"]):
                gui_max_pwm = (
                    float(gui_metadata["gui_max_pwm_us"])
                    if not np.isfinite(gui_max_pwm)
                    else max(gui_max_pwm, float(gui_metadata["gui_max_pwm_us"]))
                )

        comparison_rows.append(
            {
                "testo_row": row + 1,
                "testo_col": col + 1,
                "distance_m": measurement.distance_m,
                "testo_start_time": _format_datetime(measurement.start_time),
                "testo_end_time": _format_datetime(measurement.end_time),
                "testo_duration_s": measurement.duration_s,
                "testo_mean_ms": measurement.mean_ms,
                "testo_median_ms": measurement.median_ms,
                "testo_std_ms": measurement.std_ms,
                "gui_mean_pwm_us": gui_pwm_grid[row, col],
                "gui_mean_tach_hz": gui_tach_grid[row, col],
                "gui_mean_rpm": gui_rpm_grid[row, col],
                "gui_interval_mean_pwm_us": gui_metadata["gui_mean_pwm_us"],
                "matched_gui_log": match_path.name if match_path is not None else None,
                "matched_gui_start_time": _format_datetime(gui_metadata["start_time"]),
                "matched_gui_end_time": _format_datetime(gui_metadata["end_time"]),
                "matched_gui_duration_s": gui_metadata["duration_s"],
                "matched_gui_rows": gui_metadata["gui_log_sample_count"],
                "match_overlap_s": match_overlap_s,
                "sample_count": measurement.sample_count,
                "source_file": Path(measurement.source_path).name,
            }
        )

    testo_durations = np.asarray([m.duration_s for m in measurements], dtype=np.float64)
    finite_testo_durations = _finite(testo_durations)
    correlations = {
        "testo_vs_pwm": _pearson_from_rows(comparison_rows, "gui_mean_pwm_us"),
        "testo_vs_tach_hz": _pearson_from_rows(comparison_rows, "gui_mean_tach_hz"),
        "testo_vs_rpm": _pearson_from_rows(comparison_rows, "gui_mean_rpm"),
    }
    return TestoCorrelationResult(
        distance_m=float(distance_m),
        log_path=str(log_path) if log_path else None,
        testo_total_duration_s=float(np.sum(finite_testo_durations)) if finite_testo_durations.size else float("nan"),
        testo_mean_duration_s=float(np.mean(finite_testo_durations)) if finite_testo_durations.size else float("nan"),
        gui_matched_duration_s=gui_matched_duration_s if gui_log_sample_count else float("nan"),
        gui_mean_pwm_us=float(gui_pwm_total / gui_pwm_count) if gui_pwm_count else float("nan"),
        gui_min_pwm_us=float(gui_min_pwm),
        gui_max_pwm_us=float(gui_max_pwm),
        gui_log_sample_count=int(gui_log_sample_count),
        matched_log_count=len(matched_logs),
        measurements=measurements,
        testo_grid_ms=testo_grid,
        expanded_speed_grid_ms=_expand_measurement_grid_to_motor_grid(testo_grid),
        gui_pwm_grid=gui_pwm_grid,
        gui_tach_hz_grid=gui_tach_grid,
        gui_rpm_grid=gui_rpm_grid,
        comparison_rows=comparison_rows,
        correlations=correlations,
    )


def _make_group_summary(
    label_name: str,
    label_value: int,
    motor_ids: list[int],
    motor_means: np.ndarray,
) -> dict[str, Any]:
    values = motor_means[motor_ids]
    return {
        label_name: label_value + 1,
        "motor_count": len(motor_ids),
        "mean_hz": _stat(values, np.mean),
        "median_hz": _stat(values, np.median),
        "std_hz": _stat(values, np.std),
        "min_hz": _stat(values, np.min),
        "max_hz": _stat(values, np.max),
        "mean_rpm": _stat(values * 60.0 / TACH_PULSES_PER_REV, np.mean),
    }


def summarize_anemometer_samples(
    samples: list[AnemometerSample],
    commanded_pwm_us: float,
) -> AnemometerCalibrationResult:
    """Summarize raw telemetry and compute per-motor PWM compensation offsets."""
    if not samples:
        raise ValueError("No anemometer samples were collected")

    sample_matrix = np.vstack([np.asarray(sample.tach_hz, dtype=np.float64) for sample in samples])
    if sample_matrix.shape[1] != NUM_MOTORS:
        raise ValueError(f"Expected {NUM_MOTORS} motors, got {sample_matrix.shape[1]}")

    motor_means = np.array([_stat(sample_matrix[:, i], np.mean) for i in range(NUM_MOTORS)])
    motor_medians = np.array([_stat(sample_matrix[:, i], np.median) for i in range(NUM_MOTORS)])
    motor_stds = np.array([_stat(sample_matrix[:, i], np.std) for i in range(NUM_MOTORS)])
    motor_mins = np.array([_stat(sample_matrix[:, i], np.min) for i in range(NUM_MOTORS)])
    motor_maxes = np.array([_stat(sample_matrix[:, i], np.max) for i in range(NUM_MOTORS)])
    finite_means = _finite(motor_means)
    target_hz = float(np.mean(finite_means)) if finite_means.size else float("nan")
    target_rpm = target_hz * 60.0 / TACH_PULSES_PER_REV if np.isfinite(target_hz) else float("nan")

    suggested_pwm = np.full(NUM_MOTORS, float(commanded_pwm_us), dtype=np.float64)
    compensation_factor = np.full(NUM_MOTORS, np.nan, dtype=np.float64)
    if np.isfinite(target_hz) and target_hz > 0:
        pwm_span = max(1.0, float(commanded_pwm_us) - float(PWM_MIN))
        for motor_id, mean_hz in enumerate(motor_means):
            if np.isfinite(mean_hz) and mean_hz > 0:
                compensation_factor[motor_id] = target_hz / mean_hz
                suggested_pwm[motor_id] = PWM_MIN + pwm_span * compensation_factor[motor_id]

    suggested_pwm = np.clip(suggested_pwm, PWM_MIN, PWM_MAX)
    offset_us = suggested_pwm - float(commanded_pwm_us)

    motor_summary: list[dict[str, Any]] = []
    for motor_id in range(NUM_MOTORS):
        row, col = MOTOR_GRID_POSITION[motor_id]
        finite_count = int(np.count_nonzero(np.isfinite(sample_matrix[:, motor_id])))
        mean_hz = motor_means[motor_id]
        deviation_pct = (
            100.0 * (mean_hz - target_hz) / target_hz
            if np.isfinite(mean_hz) and np.isfinite(target_hz) and target_hz > 0
            else float("nan")
        )
        motor_summary.append(
            {
                "motor_id": motor_id + 1,
                "grid_row": row + 1,
                "grid_col": col + 1,
                "sample_count": finite_count,
                "mean_hz": mean_hz,
                "median_hz": motor_medians[motor_id],
                "std_hz": motor_stds[motor_id],
                "min_hz": motor_mins[motor_id],
                "max_hz": motor_maxes[motor_id],
                "mean_rpm": mean_hz * 60.0 / TACH_PULSES_PER_REV
                if np.isfinite(mean_hz)
                else float("nan"),
                "deviation_pct": deviation_pct,
                "compensation_factor": compensation_factor[motor_id],
                "suggested_pwm_us": suggested_pwm[motor_id],
                "offset_us": offset_us[motor_id],
            }
        )

    row_summary = [
        _make_group_summary("grid_row", row, list(WALL_MOTOR_GRID[row]), motor_means)
        for row in range(GRID_ROWS)
    ]
    column_summary = [
        _make_group_summary(
            "grid_col",
            col,
            [WALL_MOTOR_GRID[row][col] for row in range(GRID_ROWS)],
            motor_means,
        )
        for col in range(GRID_COLS)
    ]

    return AnemometerCalibrationResult(
        commanded_pwm_us=float(commanded_pwm_us),
        target_hz=target_hz,
        target_rpm=target_rpm,
        sample_count=len(samples),
        motor_summary=motor_summary,
        row_summary=row_summary,
        column_summary=column_summary,
        mean_hz_by_motor=motor_means,
        suggested_pwm_by_motor=suggested_pwm,
        offset_us_by_motor=offset_us,
        mean_hz_grid=_motor_values_to_grid(motor_means),
        suggested_pwm_grid=_motor_values_to_grid(suggested_pwm),
        offset_us_grid=_motor_values_to_grid(offset_us),
    )


def filter_grid_by_row_col(
    grid: np.ndarray,
    row_index: int | None = None,
    col_index: int | None = None,
) -> np.ndarray:
    """Return a heatmap grid with non-selected rows/columns masked as NaN."""
    filtered = np.asarray(grid, dtype=np.float64).copy()
    if row_index is not None:
        mask = np.ones(filtered.shape[0], dtype=bool)
        mask[row_index] = False
        filtered[mask, :] = np.nan
    if col_index is not None:
        mask = np.ones(filtered.shape[1], dtype=bool)
        mask[col_index] = False
        filtered[:, mask] = np.nan
    return filtered


def _excel_value(value: Any) -> Any:
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and not np.isfinite(value):
        return None
    return value


def _write_table(ws, headers: list[str], rows: list[dict[str, Any]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([_excel_value(row.get(header)) for header in headers])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)
        cell.fill = cell.fill.copy(fgColor="D9EAF7", fill_type="solid")
    _auto_size_columns(ws)


def _write_grid(ws, title: str, grid: np.ndarray, number_format: str = "0.0") -> None:
    ws["A1"] = title
    ws["A1"].font = ws["A1"].font.copy(bold=True)
    ws.append([""] + [f"Col {col + 1}" for col in range(GRID_COLS)])
    for row in range(GRID_ROWS):
        ws.append([f"Row {row + 1}"] + [_excel_value(float(grid[row, col])) for col in range(GRID_COLS)])

    for cell in ws[2]:
        cell.font = cell.font.copy(bold=True)
    for row in ws.iter_rows(min_row=3, max_row=GRID_ROWS + 2, min_col=1, max_col=GRID_COLS + 1):
        row[0].font = row[0].font.copy(bold=True)
        for cell in row[1:]:
            cell.number_format = number_format
    ws.freeze_panes = "B3"
    _auto_size_columns(ws)


def _write_any_grid(ws, title: str, grid: np.ndarray, number_format: str = "0.0") -> str:
    """Write any rectangular grid and return its data cell range."""
    grid = np.asarray(grid, dtype=np.float64)
    rows, cols = grid.shape
    ws["A1"] = title
    ws["A1"].font = ws["A1"].font.copy(bold=True)
    ws.append([""] + [f"Col {col + 1}" for col in range(cols)])
    for row in range(rows):
        ws.append([f"Row {row + 1}"] + [_excel_value(float(grid[row, col])) for col in range(cols)])

    for cell in ws[2]:
        cell.font = cell.font.copy(bold=True)
    for row in ws.iter_rows(min_row=3, max_row=rows + 2, min_col=1, max_col=cols + 1):
        row[0].font = row[0].font.copy(bold=True)
        for cell in row[1:]:
            cell.number_format = number_format
    ws.freeze_panes = "B3"
    _auto_size_columns(ws)

    from openpyxl.utils import get_column_letter

    return f"B3:{get_column_letter(cols + 1)}{rows + 2}"


def _auto_size_columns(ws) -> None:
    from openpyxl.utils import get_column_letter

    for column_cells in ws.columns:
        max_len = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = cell.value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[column_letter].width = min(max(max_len + 2, 10), 22)


def write_calibration_workbook(
    result: AnemometerCalibrationResult,
    samples: list[AnemometerSample],
    output_dir: str | Path,
) -> Path:
    """Write raw samples, summaries, heatmaps, and compensation to an .xlsx file."""
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import ColorScaleRule
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to write Excel calibration workbooks") from exc

    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    workbook_path = output_path / f"anemometer_uniform_flow_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    wb = Workbook()
    raw_ws = wb.active
    raw_ws.title = "Raw Samples"

    raw_headers = [
        "sample_index",
        "timestamp",
        "elapsed_s",
        "grid_row",
        "grid_col",
        "motor_id",
        "commanded_pwm_us",
        "tach_hz",
        "rpm",
    ]
    raw_ws.append(raw_headers)
    for sample_index, sample in enumerate(samples, start=1):
        for motor_id, tach_hz in enumerate(sample.tach_hz):
            row, col = MOTOR_GRID_POSITION[motor_id]
            rpm = tach_hz * 60.0 / TACH_PULSES_PER_REV if np.isfinite(tach_hz) else float("nan")
            raw_ws.append(
                [
                    sample_index,
                    sample.timestamp,
                    sample.elapsed_s,
                    row + 1,
                    col + 1,
                    motor_id + 1,
                    result.commanded_pwm_us,
                    _excel_value(float(tach_hz)),
                    _excel_value(float(rpm)),
                ]
            )
    raw_ws.freeze_panes = "A2"
    raw_ws.auto_filter.ref = raw_ws.dimensions
    for cell in raw_ws[1]:
        cell.font = cell.font.copy(bold=True)
        cell.fill = cell.fill.copy(fgColor="D9EAF7", fill_type="solid")
    _auto_size_columns(raw_ws)

    motor_ws = wb.create_sheet("Motor Summary")
    motor_headers = list(result.motor_summary[0].keys())
    _write_table(motor_ws, motor_headers, result.motor_summary)

    row_ws = wb.create_sheet("Row Summary")
    _write_table(row_ws, list(result.row_summary[0].keys()), result.row_summary)

    col_ws = wb.create_sheet("Column Summary")
    _write_table(col_ws, list(result.column_summary[0].keys()), result.column_summary)

    heat_ws = wb.create_sheet("Heatmap Hz")
    _write_grid(heat_ws, "Mean tach Hz by row/column", result.mean_hz_grid)
    heat_ws.conditional_formatting.add(
        f"B3:{chr(ord('A') + GRID_COLS)}{GRID_ROWS + 2}",
        ColorScaleRule(
            start_type="min",
            start_color="63BE7B",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFEB84",
            end_type="max",
            end_color="F8696B",
        ),
    )

    pwm_ws = wb.create_sheet("Suggested PWM")
    _write_grid(pwm_ws, "Suggested PWM us by row/column", result.suggested_pwm_grid, "0")
    pwm_ws.conditional_formatting.add(
        f"B3:{chr(ord('A') + GRID_COLS)}{GRID_ROWS + 2}",
        ColorScaleRule(
            start_type="min",
            start_color="63BE7B",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFEB84",
            end_type="max",
            end_color="F8696B",
        ),
    )

    offset_ws = wb.create_sheet("PWM Offsets")
    _write_grid(offset_ws, "PWM offset us by row/column", result.offset_us_grid, "0.0")
    offset_ws.conditional_formatting.add(
        f"B3:{chr(ord('A') + GRID_COLS)}{GRID_ROWS + 2}",
        ColorScaleRule(
            start_type="min",
            start_color="63BE7B",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFEB84",
            end_type="max",
            end_color="F8696B",
        ),
    )

    meta_ws = wb.create_sheet("Calibration")
    meta_rows = [
        {"field": "commanded_pwm_us", "value": result.commanded_pwm_us},
        {"field": "target_hz", "value": result.target_hz},
        {"field": "target_rpm", "value": result.target_rpm},
        {"field": "sample_count", "value": result.sample_count},
        {"field": "grid_rows", "value": GRID_ROWS},
        {"field": "grid_cols", "value": GRID_COLS},
    ]
    _write_table(meta_ws, ["field", "value"], meta_rows)

    wb.save(workbook_path)
    return workbook_path


def write_testo_correlation_workbook(
    result: TestoCorrelationResult,
    output_dir: str | Path,
) -> Path:
    """Write imported Testo points, GUI log correlation, and heatmaps to Excel."""
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import ColorScaleRule
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to write Excel correlation workbooks") from exc

    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    workbook_path = output_path / f"testo_wind_map_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    wb = Workbook()
    points_ws = wb.active
    points_ws.title = "Testo Points"
    point_headers = [
        "testo_row",
        "testo_col",
        "distance_m",
        "start_time",
        "end_time",
        "duration_s",
        "sample_count",
        "mean_ms",
        "median_ms",
        "std_ms",
        "min_ms",
        "max_ms",
        "pressure_hpa_mean",
        "source_file",
    ]
    point_rows = [
        {
            "testo_row": item.row + 1,
            "testo_col": item.col + 1,
            "distance_m": item.distance_m,
            "start_time": _format_datetime(item.start_time),
            "end_time": _format_datetime(item.end_time),
            "duration_s": item.duration_s,
            "sample_count": item.sample_count,
            "mean_ms": item.mean_ms,
            "median_ms": item.median_ms,
            "std_ms": item.std_ms,
            "min_ms": item.min_ms,
            "max_ms": item.max_ms,
            "pressure_hpa_mean": item.pressure_hpa_mean,
            "source_file": Path(item.source_path).name,
        }
        for item in sorted(result.measurements, key=lambda entry: (entry.row, entry.col))
    ]
    _write_table(points_ws, point_headers, point_rows)

    merged_ws = wb.create_sheet("Merged Blocks")
    _write_table(merged_ws, list(result.comparison_rows[0].keys()), result.comparison_rows)

    heatmaps = [
        ("Testo ms 4x4", "Mean Testo velocity m/s", result.testo_grid_ms, "0.00"),
        ("Testo ms 8x8", "Testo velocity mapped to motor blocks", result.expanded_speed_grid_ms, "0.00"),
        ("GUI PWM 4x4", "GUI mean PWM us by Testo block", result.gui_pwm_grid, "0"),
        ("GUI Tach 4x4", "GUI mean tach Hz by Testo block", result.gui_tach_hz_grid, "0.0"),
        ("GUI RPM 4x4", "GUI mean RPM by Testo block", result.gui_rpm_grid, "0"),
    ]
    for sheet_name, title, grid, number_format in heatmaps:
        ws = wb.create_sheet(sheet_name)
        cell_range = _write_any_grid(ws, title, grid, number_format)
        if np.isfinite(grid).any():
            ws.conditional_formatting.add(
                cell_range,
                ColorScaleRule(
                    start_type="min",
                    start_color="63BE7B",
                    mid_type="percentile",
                    mid_value=50,
                    mid_color="FFEB84",
                    end_type="max",
                    end_color="F8696B",
                ),
            )

    meta_ws = wb.create_sheet("Correlation")
    meta_rows = [
        {"field": "distance_m", "value": result.distance_m},
        {"field": "gui_log_selection", "value": result.log_path or "auto time match"},
        {"field": "testo_file_count", "value": len(result.measurements)},
        {"field": "testo_total_duration_s", "value": result.testo_total_duration_s},
        {"field": "testo_mean_duration_s", "value": result.testo_mean_duration_s},
        {"field": "gui_matched_duration_s", "value": result.gui_matched_duration_s},
        {"field": "gui_mean_pwm_us", "value": result.gui_mean_pwm_us},
        {"field": "gui_min_pwm_us", "value": result.gui_min_pwm_us},
        {"field": "gui_max_pwm_us", "value": result.gui_max_pwm_us},
        {"field": "gui_log_sample_count", "value": result.gui_log_sample_count},
        {"field": "matched_log_count", "value": result.matched_log_count},
        {"field": "testo_vs_pwm", "value": result.correlations["testo_vs_pwm"]},
        {"field": "testo_vs_tach_hz", "value": result.correlations["testo_vs_tach_hz"]},
        {"field": "testo_vs_rpm", "value": result.correlations["testo_vs_rpm"]},
    ]
    _write_table(meta_ws, ["field", "value"], meta_rows)

    wb.save(workbook_path)
    return workbook_path


def _write_csv_table(path: Path, headers: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({header: _excel_value(row.get(header)) for header in headers})


def _write_csv_grid(path: Path, grid: np.ndarray) -> None:
    grid = np.asarray(grid, dtype=np.float64)
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow([""] + [f"C{col + 1}" for col in range(grid.shape[1])])
        for row in range(grid.shape[0]):
            writer.writerow([f"R{row + 1}"] + [_excel_value(float(grid[row, col])) for col in range(grid.shape[1])])


def write_testo_correlation_csv_bundle(
    result: TestoCorrelationResult,
    output_dir: str | Path,
) -> Path:
    """Write the Testo/GUI merged heatmaps as a dependency-free CSV bundle."""
    bundle_dir = Path(output_dir) / f"testo_wind_map_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    bundle_dir.mkdir(parents=True, exist_ok=True)

    point_rows = [
        {
            "testo_row": item.row + 1,
            "testo_col": item.col + 1,
            "distance_m": item.distance_m,
            "start_time": _format_datetime(item.start_time),
            "end_time": _format_datetime(item.end_time),
            "duration_s": item.duration_s,
            "sample_count": item.sample_count,
            "mean_ms": item.mean_ms,
            "median_ms": item.median_ms,
            "std_ms": item.std_ms,
            "min_ms": item.min_ms,
            "max_ms": item.max_ms,
            "pressure_hpa_mean": item.pressure_hpa_mean,
            "source_file": Path(item.source_path).name,
        }
        for item in sorted(result.measurements, key=lambda entry: (entry.row, entry.col))
    ]
    _write_csv_table(bundle_dir / "testo_points.csv", list(point_rows[0].keys()), point_rows)
    _write_csv_table(bundle_dir / "merged_blocks.csv", list(result.comparison_rows[0].keys()), result.comparison_rows)

    _write_csv_grid(bundle_dir / "testo_ms_4x4.csv", result.testo_grid_ms)
    _write_csv_grid(bundle_dir / "testo_ms_8x8_motor_blocks.csv", result.expanded_speed_grid_ms)
    _write_csv_grid(bundle_dir / "gui_pwm_4x4.csv", result.gui_pwm_grid)
    _write_csv_grid(bundle_dir / "gui_tach_hz_4x4.csv", result.gui_tach_hz_grid)
    _write_csv_grid(bundle_dir / "gui_rpm_4x4.csv", result.gui_rpm_grid)

    metadata_rows = [
        {"field": "distance_m", "value": result.distance_m},
        {"field": "gui_log_selection", "value": result.log_path or "auto time match"},
        {"field": "testo_file_count", "value": len(result.measurements)},
        {"field": "testo_total_duration_s", "value": result.testo_total_duration_s},
        {"field": "testo_mean_duration_s", "value": result.testo_mean_duration_s},
        {"field": "gui_matched_duration_s", "value": result.gui_matched_duration_s},
        {"field": "gui_mean_pwm_us", "value": result.gui_mean_pwm_us},
        {"field": "gui_min_pwm_us", "value": result.gui_min_pwm_us},
        {"field": "gui_max_pwm_us", "value": result.gui_max_pwm_us},
        {"field": "gui_log_sample_count", "value": result.gui_log_sample_count},
        {"field": "matched_log_count", "value": result.matched_log_count},
        {"field": "testo_vs_pwm", "value": result.correlations["testo_vs_pwm"]},
        {"field": "testo_vs_tach_hz", "value": result.correlations["testo_vs_tach_hz"]},
        {"field": "testo_vs_rpm", "value": result.correlations["testo_vs_rpm"]},
    ]
    _write_csv_table(bundle_dir / "correlation.csv", ["field", "value"], metadata_rows)
    return bundle_dir
