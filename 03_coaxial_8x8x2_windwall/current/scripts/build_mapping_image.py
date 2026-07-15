#!/usr/bin/env python3
"""Render a precise 8x8x2 controller mapping image."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from config import GRID_COLS, GRID_ROWS  # noqa: E402
from coaxial_windwall.model import CoaxialAddress, controller_host_indices, address_from_motor_index  # noqa: E402


DOCS_OUTPUT = PROJECT_ROOT / "docs" / "coaxial_8x8x2_pico_mapping.png"
DESKTOP_OUTPUT = Path.home() / "Desktop" / "coaxial_8x8x2_pico_mapping.png"
PROPOSED_DESKTOP_OUTPUT = Path.home() / "Desktop" / "coaxial_8x8x2_pico_mapping_proposed.png"

FONT_REGULAR = "/usr/share/fonts/opentype/urw-base35/NimbusSans-Regular.otf"
FONT_BOLD = "/usr/share/fonts/opentype/urw-base35/NimbusSans-Bold.otf"

INK = (20, 24, 31)
MUTED = (91, 101, 115)
GRID = (189, 196, 207)
BG = (248, 250, 252)
PANEL = (255, 255, 255)
OLD = (37, 99, 235)
INFILL = (217, 119, 6)
HEADER = (17, 24, 39)
WHITE = (255, 255, 255)

CONTROLLER_COLORS = (
    (30, 100, 220),   # C01 blue
    (0, 150, 95),     # C02 emerald
    (210, 45, 45),    # C03 red
    (130, 70, 200),   # C04 purple
    (230, 125, 30),   # C05 orange
    (0, 145, 170),    # C06 cyan
    (215, 60, 135),   # C07 pink
    (105, 150, 25),   # C08 olive
    (180, 140, 0),    # C09 gold
    (130, 25, 155),   # C10 violet
    (0, 125, 120),    # C11 teal
    (150, 75, 35),    # C12 brown
    (60, 80, 200),    # C13 indigo
    (230, 70, 80),    # C14 coral
    (60, 125, 55),    # C15 forest
    (80, 105, 125),   # C16 slate
)

CONTROLLER_FILLS = (
    (221, 235, 255),
    (214, 248, 231),
    (255, 226, 226),
    (240, 229, 255),
    (255, 237, 213),
    (207, 250, 254),
    (252, 231, 243),
    (236, 252, 203),
    (254, 249, 195),
    (250, 232, 255),
    (204, 251, 241),
    (245, 232, 220),
    (224, 231, 255),
    (255, 228, 230),
    (220, 252, 231),
    (226, 232, 240),
)


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REGULAR, size)


def text_center(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    text: str,
    fill: tuple[int, int, int],
    image_font: ImageFont.FreeTypeFont,
) -> None:
    bbox = draw.textbbox((0, 0), text, font=image_font)
    w = bbox[2] - bbox[0]
    h = bbox[3] - bbox[1]
    x0, y0, x1, y1 = xy
    draw.text(
        (x0 + (x1 - x0 - w) / 2, y0 + (y1 - y0 - h) / 2 - 1),
        text,
        fill=fill,
        font=image_font,
    )


def controller_role(controller_label: str) -> str:
    controller_index = int(controller_label[1:]) - 1
    return "old harness" if controller_index < 8 else "new infill"


def controller_index_from_label(controller_label: str) -> int:
    return int(controller_label[1:]) - 1


def controller_color(controller_label: str) -> tuple[int, int, int]:
    return CONTROLLER_COLORS[controller_index_from_label(controller_label)]


def controller_fill(controller_label: str) -> tuple[int, int, int]:
    return CONTROLLER_FILLS[controller_index_from_label(controller_label)]


def draw_cell(
    draw: ImageDraw.ImageDraw,
    x: int,
    y: int,
    w: int,
    h: int,
    row: int,
    col: int,
    pixel_number: int | None = None,
) -> None:
    front = CoaxialAddress(row=row, col=col, layer=0)
    back = CoaxialAddress(row=row, col=col, layer=1)
    top_h = 30
    half = (h - top_h) // 2

    border_color = controller_color(front.controller_label)
    front_color = controller_color(front.controller_label)
    back_color = controller_color(back.controller_label)
    front_fill = controller_fill(front.controller_label)
    back_fill = controller_fill(back.controller_label)
    draw.rounded_rectangle((x, y, x + w, y + h), radius=10, fill=PANEL, outline=border_color, width=4)

    draw.rectangle((x + 4, y + top_h, x + w - 4, y + top_h + half), fill=front_fill)
    draw.rectangle((x + 4, y + top_h + half, x + w - 4, y + h - 4), fill=back_fill)
    draw.rectangle((x + 4, y + top_h, x + 13, y + top_h + half), fill=front_color)
    draw.rectangle((x + 4, y + top_h + half, x + 13, y + h - 4), fill=back_color)
    draw.line((x + 4, y + top_h + half, x + w - 4, y + top_h + half), fill=GRID, width=2)

    pair_number = front.pixel_number if pixel_number is None else pixel_number
    pair_label = f"P{pair_number:02d}"
    front_label = f"F{pair_number:02d}"
    back_label = f"B{pair_number:02d}"
    text_center(draw, (x, y + 1, x + w, y + top_h), f"{pair_label}  R{row + 1}C{col + 1}", INK, font(18, True))

    draw.rounded_rectangle((x + 18, y + top_h + 8, x + 57, y + top_h + 34), radius=5, fill=front_color)
    text_center(draw, (x + 18, y + top_h + 8, x + 57, y + top_h + 34), "F", WHITE, font(16, True))
    draw.text((x + 64, y + top_h + 10), front_label, fill=INK, font=font(20, True))
    draw.text(
        (x + 112, y + top_h + 12),
        f"{front.controller_label}-{front.controller_channel_label}",
        fill=front_color,
        font=font(17, True),
    )

    draw.rounded_rectangle((x + 18, y + top_h + half + 8, x + 57, y + top_h + half + 34), radius=5, fill=back_color)
    text_center(draw, (x + 18, y + top_h + half + 8, x + 57, y + top_h + half + 34), "B", WHITE, font(16, True))
    draw.text((x + 64, y + top_h + half + 10), back_label, fill=INK, font=font(20, True))
    draw.text(
        (x + 112, y + top_h + half + 12),
        f"{back.controller_label}-{back.controller_channel_label}",
        fill=back_color,
        font=font(17, True),
    )


def controller_summary(
    controller_index: int,
    pixel_numbers: list[list[int]] | None = None,
) -> list[str]:
    parts = []
    for channel_index, motor_index in enumerate(controller_host_indices(controller_index), start=1):
        address = address_from_motor_index(motor_index)
        if pixel_numbers is None:
            label = address.label
        else:
            label = f"{address.layer_label}{pixel_numbers[address.row][address.col]:02d}"
        parts.append(f"CH{channel_index}:{label}")
    return parts


def draw_legend(
    draw: ImageDraw.ImageDraw,
    x: int,
    y: int,
    w: int,
    pixel_numbers: list[list[int]] | None = None,
) -> None:
    draw.rounded_rectangle((x, y, x + w, y + 1170), radius=12, fill=PANEL, outline=GRID, width=2)
    draw.text((x + 24, y + 22), "Pico / Controller Legend", fill=INK, font=font(30, True))

    draw.rounded_rectangle((x + 24, y + 72, x + 190, y + 108), radius=8, fill=(245, 247, 250), outline=OLD, width=2)
    draw.text((x + 38, y + 79), "C01-C08 old harness", fill=OLD, font=font(17, True))
    draw.rounded_rectangle((x + 210, y + 72, x + 386, y + 108), radius=8, fill=(245, 247, 250), outline=INFILL, width=2)
    draw.text((x + 224, y + 79), "C09-C16 new infill", fill=INFILL, font=font(17, True))

    draw.text((x + 24, y + 130), "Each grid cell shows:", fill=INK, font=font(18, True))
    draw.text((x + 42, y + 158), "color: owning Pico/controller C01-C16", fill=MUTED, font=font(16))
    draw.text((x + 42, y + 182), "F/B badge: front or back motor", fill=MUTED, font=font(16))

    start_y = y + 225
    row_h = 56
    for controller_index in range(16):
        row = controller_index % 8
        col = controller_index // 8
        cx = x + 24 + col * 258
        cy = start_y + row * row_h
        label = f"C{controller_index + 1:02d}"
        role = "old" if controller_index < 8 else "infill"
        color = CONTROLLER_COLORS[controller_index]
        fill = CONTROLLER_FILLS[controller_index]
        draw.rounded_rectangle((cx, cy, cx + 238, cy + 46), radius=8, fill=fill, outline=color, width=3)
        draw.rectangle((cx, cy, cx + 12, cy + 46), fill=color)
        draw.text((cx + 20, cy + 6), f"{label} {role}", fill=color, font=font(17, True))
        pairs = []
        for motor_index in controller_host_indices(controller_index)[::2]:
            address = address_from_motor_index(motor_index)
            if pixel_numbers is None:
                pairs.append(address.pair_label)
            else:
                pairs.append(f"P{pixel_numbers[address.row][address.col]:02d}")
        draw.text((cx + 20, cy + 27), ", ".join(pairs), fill=INK, font=font(14, True))

    draw.text((x + 24, y + 704), "Detailed channel order", fill=INK, font=font(22, True))
    detail_y = y + 740
    for controller_index in range(16):
        label = f"C{controller_index + 1:02d}"
        color = CONTROLLER_COLORS[controller_index]
        left = x + 24 if controller_index < 8 else x + 284
        line_y = detail_y + (controller_index % 8) * 48
        draw.text((left, line_y), label, fill=color, font=font(16, True))
        chunks = controller_summary(controller_index, pixel_numbers)
        draw.text((left + 42, line_y), " ".join(chunks[:4]), fill=INK, font=font(13))
        draw.text((left + 42, line_y + 18), " ".join(chunks[4:]), fill=INK, font=font(13))


def load_pixel_numbers(workbook_path: Path) -> list[list[int]]:
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    sheet = workbook.active
    values = [
        [int(sheet.cell(row=row + 2, column=col + 2).value) for col in range(GRID_COLS)]
        for row in range(GRID_ROWS)
    ]
    flattened = [number for row in values for number in row]
    if sorted(flattened) != list(range(1, GRID_ROWS * GRID_COLS + 1)):
        raise ValueError("workbook grid must contain each motor number 1-64 exactly once")
    return values


def build_image(
    path: Path,
    pixel_numbers: list[list[int]] | None = None,
    source_name: str | None = None,
) -> None:
    width = 2400
    height = 1500
    margin = 60
    title_h = 120
    cell_w = 200
    cell_h = 145
    gap = 10
    grid_x = margin
    grid_y = title_h + 80

    image = Image.new("RGB", (width, height), BG)
    draw = ImageDraw.Draw(image)

    draw.rounded_rectangle((0, 0, width, 120), radius=0, fill=HEADER)
    title = "Proposed Coaxial 8x8x2 Pico Mapping" if pixel_numbers else "Coaxial 8x8x2 Pico Mapping"
    draw.text((margin, 30), title, fill=WHITE, font=font(42, True))
    draw.text(
        (margin, 78),
        (
            "Motor numbers from Book.xlsx; Pico positions and channels unchanged from the current coaxial layout."
            if pixel_numbers
            else "Original 8x8 Pico-block numbering. C01-C08 retain old harness positions; C09-C16 are infill."
        ),
        fill=(203, 213, 225),
        font=font(22),
    )

    for col in range(GRID_COLS):
        text_center(
            draw,
            (grid_x + col * (cell_w + gap), grid_y - 38, grid_x + col * (cell_w + gap) + cell_w, grid_y - 6),
            f"C{col + 1}",
            MUTED,
            font(18, True),
        )
    for row in range(GRID_ROWS):
        text_center(
            draw,
            (grid_x - 46, grid_y + row * (cell_h + gap), grid_x - 8, grid_y + row * (cell_h + gap) + cell_h),
            f"R{row + 1}",
            MUTED,
            font(18, True),
        )

    for row in range(GRID_ROWS):
        for col in range(GRID_COLS):
            x = grid_x + col * (cell_w + gap)
            y = grid_y + row * (cell_h + gap)
            number = pixel_numbers[row][col] if pixel_numbers else None
            draw_cell(draw, x, y, cell_w, cell_h, row, col, number)

    legend_x = grid_x + GRID_COLS * (cell_w + gap) + 60
    draw_legend(draw, legend_x, grid_y - 46, width - legend_x - margin, pixel_numbers)

    draw.text(
        (margin, height - 42),
        (
            f"PROPOSED numbering from {source_name}; Pico/controller positions retained from current mapping"
            if pixel_numbers
            else "Generated from coaxial_windwall.model.CONTROLLER_HOST_INDICES and docs/MOTOR_CONTROLLER_MAPPING.md"
        ),
        fill=MUTED,
        font=font(18),
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DOCS_OUTPUT)
    parser.add_argument("--desktop-copy", action="store_true")
    parser.add_argument("--numbering-workbook", type=Path)
    args = parser.parse_args()

    pixel_numbers = load_pixel_numbers(args.numbering_workbook) if args.numbering_workbook else None
    output = PROPOSED_DESKTOP_OUTPUT if args.numbering_workbook and args.output == DOCS_OUTPUT else args.output
    build_image(
        output,
        pixel_numbers=pixel_numbers,
        source_name=args.numbering_workbook.name if args.numbering_workbook else None,
    )
    print(output)
    if args.desktop_copy:
        desktop_output = PROPOSED_DESKTOP_OUTPUT if pixel_numbers else DESKTOP_OUTPUT
        build_image(
            desktop_output,
            pixel_numbers=pixel_numbers,
            source_name=args.numbering_workbook.name if args.numbering_workbook else None,
        )
        print(desktop_output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
